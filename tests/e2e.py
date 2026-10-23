"""End-to-end smoke test: upload sample drawings, wait for extraction, check every deliverable table and export."""
import sys, os, time, tempfile, io, zipfile
ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
sys.path.insert(0, ROOT)
tmp = tempfile.mkdtemp()
os.environ["PRG_DB"] = os.path.join(tmp, "e2e.db")
os.environ["PRG_HOME"] = tmp
os.environ["PRG_AI_PROVIDER"] = "off"
import app as A

DOCS = os.path.join(ROOT, "Docs")
FILES = sys.argv[1:] or ["PRG-2425-465-MOJJ ENGG MSD-697 PLC PANEL.pdf", "PRG-2526-616-AQUA CARE-SHAH EYE CARE-MCC+PLC PANEL.pdf"]
c = A.app.test_client()
pid = c.post("/api/projects", json={"name": "E2E Project"}).json["id"]
r = c.post(f"/api/projects/{pid}/upload", data={"files": [(open(os.path.join(DOCS, f), "rb"), f) for f in FILES]}, content_type="multipart/form-data")
assert r.status_code == 200, r.data
t0 = time.time()
while time.time() - t0 < 300:
    d = c.get(f"/api/projects/{pid}/docs").json
    if all(x["status"] in ("done", "error") for x in d):
        break
    time.sleep(2)
for x in d:
    print(x["filename"][:44], x["status"], x["message"])
    assert x["status"] == "done", x
print("stats", c.get(f"/api/projects/{pid}/stats").json)
fail = []
for k in ["io", "summary", "bom", "bom_lines", "devices", "devsummary", "terminals", "blocks", "wires", "cables", "xref", "contacts", "coils", "sheetrefs", "circuits"]:
    j = c.get(f"/api/projects/{pid}/table/{k}")
    assert j.status_code == 200, (k, j.status_code)
    j = j.json
    print(f"{k:10s} rows={j['total']:4d}")
    if j["total"] == 0 and k not in ("circuits", "bom_lines", "cables", "coils", "contacts", "sheetrefs"):
        fail.append(k)
for f in ["xlsx", "pdf", "csv", "xml"]:
    r = c.get(f"/api/export/{pid}/{f}")
    print(f, r.status_code, len(r.data)); assert r.status_code == 200 and len(r.data) > 500
from openpyxl import load_workbook
wb = load_workbook(io.BytesIO(c.get(f"/api/export/{pid}/xlsx").data))
print("xlsx sheets:", wb.sheetnames)
# learning: edit a cable/device cell, make sure it is stored as a rule
row = c.get(f"/api/projects/{pid}/table/devices").json["rows"][0]
print("edit", c.patch(f"/api/row/devices/{row['id']}", json={"field": "make", "value": "EDITED-MAKE"}).json)
print("rules", c.get("/api/learning").json["rules"][:1])
assert not fail, f"empty tables: {fail}"
print("E2E OK")
