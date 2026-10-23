"""Datasets + exporters: XLSX (formatted), CSV, PDF, XML."""
import io, csv, zipfile, re
from collections import OrderedDict, defaultdict
from xml.sax.saxutils import escape
from . import db

IO_COLS = ["Rack", "Slot", "Module Type", "DI", "DO", "AI", "AO", "RTD", "Thermocouple", "Spare I/O", "Signal Description", "Tag Name",
           "Signal Type", "Field Device Name", "Cable Number", "Terminal Number", "Remarks", "Sheet"]


def datasets(pid):
    D = OrderedDict()
    io = db.q("SELECT * FROM io_rows WHERE project_id=? ORDER BY doc_id, CAST(rack AS INT), CAST(slot AS INT), CAST(channel AS INT)", (pid,))
    rows = []
    for r in io:
        t = r["module_type"]
        mark = lambda k: r["plc_address"] if t == k and not r["spare"] else ""
        rows.append([r["rack"], r["slot"], f"{r['module_name']} ({r['module_model']})" if r["module_model"] else r["module_name"],
                     mark("DI"), mark("DO"), mark("AI"), mark("AO"), mark("RTD"), mark("TC"),
                     (r["plc_address"] or "SPARE") if r["spare"] else "", r["description"], r["tag"], r["signal_type"],
                     r["field_device"], r["cable_no"], r["terminal"], r["remarks"], r["sheet"]])
    D["I-O List"] = (IO_COLS, rows)

    mods = db.q("SELECT * FROM modules WHERE project_id=? ORDER BY doc_id, CAST(rack AS INT), CAST(slot AS INT)", (pid,))
    summ = []
    for m in mods:
        rr = [r for r in io if r["doc_id"] == m["doc_id"] and r["rack"] == m["rack"] and r["slot"] == m["slot"] and r["module_type"] == m["module_type"]]
        sp = sum(1 for r in rr if r["spare"])
        summ.append([m["rack"], m["slot"], m["label"], m["model"], m["module_type"], len(rr), len(rr) - sp, sp,
                     f"{(100*sp/len(rr)):.0f}%" if rr else "", "derived from drawing order" if m["slot_derived"] == 1 else ""])
    D["I-O Summary"] = (["Rack", "Slot", "Module", "Model", "Type", "Channels", "Used", "Spare", "Spare %", "Remarks"], summ)

    w = db.q("SELECT * FROM wires WHERE project_id=? ORDER BY CAST(sheet AS INT), wire_no", (pid,))
    D["Ferrule List"] = (["Ferrule No", "Wire No", "From Terminal", "To Terminal", "Device Reference", "Cable Reference", "Core No", "Sheet", "Source"],
                         [[x["ferrule_no"], x["wire_no"], x["from_terminal"], x["to_terminal"], x["device_ref"], x["cable_ref"], x["core_no"], x["sheet"], x["source"]] for x in w])

    D["BOM"] = bom_dataset(pid)

    alld = db.q("SELECT * FROM devices WHERE project_id=? ORDER BY dev_class, tag", (pid,))
    D["Device List"] = (["Device Class", "Tag", "Description", "Make", "Part No", "Rating", "Sheet", "Source"],
                        [[d["dev_class"], d["tag"], d["descr"], d["make"], d["part_no"], d["rating"], d["sheet"], d["source"]] for d in alld])
    D["Device Summary"] = device_summary(pid, alld)

    t = db.q("SELECT * FROM terminals WHERE project_id=? ORDER BY strip, CAST(terminal_no AS INT)", (pid,))
    D["Terminal Schedule"] = (["Terminal Strip", "Terminal No", "Terminal Type", "Incoming Wire", "Outgoing Wire", "Cable No", "Core No", "Signal Tag", "PLC Address", "Spare", "Sheet"],
                              [[x["strip"], x["terminal_no"], x["terminal_type"], x["incoming_wire"], x["outgoing_wire"], x["cable_no"], x["core_no"], x["signal_tag"], x["plc_address"], x["spare"], x["sheet"]] for x in t])
    strips = OrderedDict()
    for x in t:
        g = strips.setdefault(x["strip"], dict(n=0, spare=0, sheet=x["sheet"], types=set()))
        g["n"] += 1; g["spare"] += 1 if x["spare"] else 0; g["types"].add(x["terminal_type"])
    D["Terminal Blocks"] = (["Terminal Strip", "Total Terminals", "Used", "Spare Terminals", "Spare %", "Terminal Type(s)", "Sheet"],
                            [[k, g["n"], g["n"] - g["spare"], g["spare"], f"{100*g['spare']/g['n']:.0f}%", ", ".join(sorted(g["types"])), g["sheet"]] for k, g in strips.items()])
    c = db.q("SELECT * FROM cables WHERE project_id=? ORDER BY id", (pid,))
    D["Cable Schedule"] = (["Cable No", "Cable Type", "Core Details", "Source Equipment", "Destination Equipment", "Length (m)", "Remarks", "Sheet"],
                           [[x["cable_no"], x["cable_type"], x["cores"], x["source_equipment"], x["dest_equipment"], x["length_m"], x["remarks"], x["sheet"]] for x in c])
    D.update(xref_datasets(pid))
    sy = db.q("SELECT * FROM symbols WHERE project_id=? ORDER BY CAST(sheet AS INT), CAST(col AS INT), tag", (pid,))
    D["Symbols & Pins"] = (["Sheet", "Column", "Tag", "Symbol Role", "Device Class", "Pins", "Rating"], [[x["sheet"], x["col"], x["tag"], x["role"], x["dev_class"], x["pins"], x["rating"]] for x in sy])
    an = db.q("SELECT * FROM annotations WHERE project_id=? ORDER BY CAST(sheet AS INT), id", (pid,))
    D["Sheet Annotations"] = (["Sheet", "Text", "Nearest Device", "Kind"], [[x["sheet"], x["text"], x["near_tag"], x["kind"]] for x in an])
    circs = db.q("SELECT * FROM circuits WHERE project_id=? ORDER BY CAST(sheet AS INT), id", (pid,))
    if circs:
        D["Feeders - Circuits"] = (
            ["Feeder No", "Description", "Rating (kW)", "Starter Type", "MCB", "Contactor", "OLR", "Cable Size", "Terminal", "Controls", "Sheet"],
            [[x["feeder_no"], x["description"], x["rating_kw"], x["starter_type"], x["mcb"], x["contactor"], x["olr"], x["cable_size"], x["terminal"], x["controls"], x["sheet"]] for x in circs]
        )

    # Drawing Register — one row per page, with full classification metadata
    sh = db.q(
        "SELECT s.sheet, s.label, s.title, s.draw_type_label, s.drawing_no, s.revision, "
        "s.kind, s.scanned, d.filename "
        "FROM sheets s JOIN documents d ON d.id = s.doc_id "
        "WHERE s.project_id=? ORDER BY s.doc_id, s.sheet", (pid,)
    )
    D["Drawing Register"] = (
        ["Page", "Sheet No", "Drawing Title", "Drawing Type", "Drawing No.", "Revision", "Kind", "Scanned / Vector", "Source File"],
        [[r["sheet"], r["label"] or "", r["title"] or "", r["draw_type_label"] or r["kind"] or "",
          r["drawing_no"] or "", r["revision"] or "",
          r["kind"] or "", "Scanned (OCR)" if r["scanned"] else "Vector", r["filename"] or ""] for r in sh]
    )

    # Panel Specifications & GA Details
    ps = db.q("SELECT sheet, category, property_name, property_value, unit, raw_text FROM panel_specs WHERE project_id=? ORDER BY CAST(sheet AS INT), id", (pid,))
    if ps:
        D["Panel Specifications"] = (
            ["Sheet", "Category", "Property", "Specification / Value", "Unit", "Drawing Reference"],
            [[r["sheet"], r["category"], r["property_name"], r["property_value"], r["unit"] or "", r["raw_text"] or ""] for r in ps]
        )

    return D


def _num(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return 1


RATED_CLASSES = {"MCB": "MCB", "MCCB": "MCCB", "Contactor": "Power contactor", "Overload Relay / Motor Protector": "Overload relay",
                 "Relay": "Control relay", "Power Supply": "Power supply (SMPS)", "Lamp": "Indication lamp"}
TAG_LIKE = re.compile(r"^\d+[A-Z]?(?:/\d+[A-Z]?)?-[A-Z]{1,5}\d*$")     # 9F-T1 etc. are device tags, not part numbers


def bom_dataset(pid):
    """Bill of material: printed BOM sheets; otherwise built from the drawings - every rated device recognised on the schematics
    (MCB / MCCB / contactor / overload ... grouped by rating), PLC modules (one per module), catalogue part numbers and cables.
    Duplicates are merged and quantities totalled; make / part number stay blank where the drawing does not print them."""
    agg = OrderedDict()

    def add(key, qty, sheets=(), tags=(), **kw):
        a = agg.setdefault(key, dict(descr="", make="", part="", rating="", cat="", panel="", sheets=set(), qty=0, lines=0, uom="", tags=[], src=""))
        for k, v in kw.items():
            if v and not a.get(k):
                a[k] = v
        a["qty"] += qty; a["lines"] += 1; a["sheets"].update(_num(x) for x in sheets if str(x).strip()); a["tags"] += [t for t in tags if t]
    printed_docs = {r["doc_id"] for r in db.q("SELECT DISTINCT doc_id FROM bom_items WHERE project_id=?", (pid,))}
    for b in db.q("SELECT * FROM bom_items WHERE project_id=? ORDER BY CAST(sheet AS INT), CAST(item_no AS INT)", (pid,)):
        add(("bom", b["make"].upper(), (b["part_no"] or b["descr"]).upper(), b["panel"].upper()), _num(b["qty"]), [b["sheet"]], descr=b["descr"], make=b["make"],
            part=b["part_no"], rating=b["rating"], cat=b["catalogue_no"], panel=b["panel"], uom=b["uom"], src="BOM sheet")
    printed = {k[2] for k in agg}
    # PLC modules: one per module (its part number is printed on every sheet of the module)
    mods = OrderedDict()
    for m in db.q("SELECT * FROM modules WHERE project_id=? ORDER BY doc_id, CAST(slot AS INT)", (pid,)):
        if m["doc_id"] in printed_docs:
            continue
        mods.setdefault((m["doc_id"], m["rack"], m["slot"], m["label"]), m)
    plc_parts = set()
    for m in mods.values():
        part = re.sub(r"^MLFB-", "", m["model"] or "", flags=re.I)
        plc_parts.add(part.upper())
        kind = "PLC CPU (with on-board I/O)" if m["label"].startswith("CPU") else f"{ {'DI': 'Digital input', 'DO': 'Digital output', 'AI': 'Analog input', 'AO': 'Analog output'}.get(m['module_type'], m['module_type'])} module, {m['channels']} ch"
        add(("plc", part.upper() or m["label"]), 1, m["sheets"].split(","), descr=kind, make="Delta" if part.upper().startswith("DVP") else "", part=part, cat=part, src="Drawing (PLC modules)")
    # rated devices recognised on the schematics: one per device tag, grouped by type + rating
    best = OrderedDict()
    for sy in db.q("SELECT s.* FROM symbols s WHERE s.project_id=? ORDER BY CAST(s.sheet AS INT)", (pid,)):
        if sy["doc_id"] in printed_docs or sy["dev_class"] not in RATED_CLASSES:
            continue
        e = best.setdefault((sy["doc_id"], sy["tag"]), dict(cls=sy["dev_class"], ratings=[], sheets=set()))
        if sy["rating"]:
            e["ratings"].append(sy["rating"])
        e["sheets"].add(sy["sheet"])
    # the single-line diagram's feeder table states each feeder device's rating explicitly - it wins over symbol proximity
    sld = {}
    for c in db.q("SELECT * FROM circuits WHERE project_id=?", (pid,)):
        for f in ("mcb", "contactor", "olr"):
            for m in re.finditer(r"(\d+F-[A-Z]+\d?)\s+((?:\d[PN]\s*,\s*)?\d+(?:\.\d+)?(?:-\d+(?:\.\d+)?)?A)", c[f] or ""):
                sld[(c["doc_id"], m.group(1))] = m.group(2).replace(" ", "")
    for (doc, tag), e in best.items():
        r_sld = sld.get((doc, tag), "")
        if re.search(r"\dP,|\d-\d", r_sld) or (r_sld and not e["ratings"]):   # fully specified (poles / OLR range) or nothing better
            e["ratings"] = [r_sld]
        if not e["ratings"]:
            continue
        rating = max(set(e["ratings"]), key=e["ratings"].count)
        add(("rated", e["cls"], rating.upper().replace(" ", "")), 1, e["sheets"], [tag], descr=f"{RATED_CLASSES[e['cls']]} {rating}", rating=rating, src="Drawing (ratings)")
    # catalogue part numbers printed on schematics (meters, HMI, ...)
    for d in db.q("SELECT * FROM devices WHERE project_id=? AND part_no<>'' AND source NOT IN ('bom','symbol') ORDER BY make, part_no", (pid,)):
        pn = d["part_no"].upper()
        if pn in printed or TAG_LIKE.match(pn) or (d["doc_id"] not in printed_docs and (pn in plc_parts or d["dev_class"] == "PLC Module")):
            continue
        add(("part", d["make"].upper(), pn), 1, [d["sheet"]], [d["tag"]], descr=d["descr"] or d["dev_class"], make=d["make"], part=d["part_no"], rating=d["rating"],
            cat=d["catalogue_no"] or d["part_no"], panel=d["panel"], src="Schematic (part no.)")
    # power / control cables by size (quantity = number of cable runs; length is not on the drawings)
    for c in db.q("SELECT * FROM cables WHERE project_id=?", (pid,)):
        if c["doc_id"] in printed_docs:
            continue
        m = re.search(r"(\d+(?:\.\d+)?)\s*sq\.?mm", (c["cores"] or "") + " " + (c["cable_type"] or ""), re.I)
        if m:
            add(("cable", m.group(1)), 1, [c["sheet"]], [c["cable_no"]], descr=f"Cable {m.group(1)} sq.mm",
                rating=f"{m.group(1)} sq.mm", uom="Runs", src="Cable schedule")
    cols = ["Part Description", "Manufacturer", "Part Number", "Quantity", "UoM", "Rating", "Catalogue Number", "Panel Reference", "Sheet Number", "Tags", "Duplicate Lines Merged", "Source"]
    order = {"bom": 0, "plc": 1, "rated": 2, "part": 3, "cable": 4}
    items = sorted(agg.items(), key=lambda kv: (order[kv[0][0]], kv[1]["descr"]))
    return cols, [[a["descr"], a["make"], a["part"], a["qty"], a["uom"] or "Nos.", a["rating"], a["cat"], a["panel"], ", ".join(str(x) for x in sorted(a["sheets"])),
                   ", ".join(sorted(set(a["tags"]), key=lambda t: (len(t), t))[:40]), a["lines"] - 1, a["src"]] for _, a in items]


def device_summary(pid, alld):
    want = ["Push Button", "Selector Switch", "Lamp", "Relay", "Contactor", "MCB", "MCCB", "PLC Module", "Power Supply", "HMI", "VFD",
            "Transmitter / Sensor", "Switch / Limit / Proximity", "Terminal Block"]
    cnt = OrderedDict((k, 0) for k in want)
    for d in alld:
        if d["source"] == "bom":
            m = re.search(r"\(qty (\d+)\)", d["descr"] or "")
            cnt[d["dev_class"]] = cnt.get(d["dev_class"], 0) + (int(m.group(1)) if m else 1)
        else:
            cnt[d["dev_class"]] = cnt.get(d["dev_class"], 0) + 1
    return ["Device Class", "Count (tagged + BOM qty)"], [[k, v] for k, v in cnt.items() if v or k in want]


def xref_datasets(pid):
    xr = db.q("SELECT * FROM xrefs WHERE project_id=? ORDER BY CAST(sheet AS INT), id", (pid,))
    out = OrderedDict()
    dev = [x for x in xr if x["direction"] == "device"]
    from .devices import classify
    out["Device Cross Reference"] = (["Tag", "Device Class", "First Sheet", "Sheets (occurrences)", "Total Occurrences"],
                                     [[x["label"], classify(x["label"]), x["sheet"], x["text"].split("; ", 1)[-1], x["target_col"]] for x in dev])
    rows = []
    for x in xr:
        if x["direction"] in ("coil", "contact"):
            parts = [t.strip() for t in x["text"].split(";")]
            role = parts[0]
            kind = next((t for t in parts[1:] if t in ("NO", "NC")), "")
            pins = next((t[5:] for t in parts if t.startswith("pins=")), "")
            refs = [r for r in next((t[5:] for t in parts if t.startswith("refs=")), "").split(",") if r and r != "-"]
            rows.append(dict(tag=x["label"], role=role, kind=kind, pins=pins, refs=refs, sheet=_num(x["sheet"]), col=_num(x["target_col"])))
    # printed references use the drawing's own sheet numbering (0-based on EPLAN exports): find the offset that makes them land on symbols
    at = {(r["tag"], r["sheet"], r["col"]) for r in rows}
    votes = {0: 0, 1: 0}
    for r in rows:
        for ref in r["refs"]:
            sh, _, c = ref.partition(".")
            for d in (0, 1):
                if (r["tag"], _num(sh) + d, _num(c)) in at:
                    votes[d] += 1
    off = 1 if votes[1] >= votes[0] else 0
    out["Contact Cross Reference"] = (["Tag", "Role", "Contact Type", "Pins", "Sheet", "Column", "Printed Cross-Refs (sheet.col -> page)"],
        [[r["tag"], r["role"], r["kind"], r["pins"], r["sheet"], r["col"], ", ".join(f"{x} -> p{_num(x.split('.')[0]) + off}" for x in r["refs"])] for r in rows])
    cm = []
    by = OrderedDict()
    for r in rows:
        by.setdefault(r["tag"], []).append(r)
    for tag, rs in by.items():
        coils = [r for r in rs if r["role"] == "coil"]
        cons = [r for r in rs if r["role"] != "coil"]
        printed = sorted({(_num(x.split(".")[0]) + off, _num(x.split(".")[1])) for r in rs for x in r["refs"]})
        if not coils and not cons:
            continue
        cm.append([tag, classify(tag), "; ".join(f"p{c['sheet']}.{c['col']}" for c in coils),
                   "; ".join(f"p{c['sheet']}.{c['col']} {c['role']}{' ' + c['kind'] if c['kind'] else ''} [{c['pins']}]" for c in cons),
                   ", ".join(f"p{a}.{b}" for a, b in printed), len(cons),
                   "coil and contacts found on the drawings" if coils and cons else ("coil not found on the drawings" if not coils else "no separate contacts found")])
    out["Coil-Contact Mapping"] = (["Device Tag", "Class", "Coil Location (page.col)", "Contacts (page.col role type [pins])", "Printed Refs (page.col)", "Contact Count", "Remarks"], cm)
    sig = [x for x in xr if x["direction"] in ("to", "from")]
    out["Signal Links"] = (["Sheet", "Direction", "Link Text", "Target Sheet", "Target Column", "Net / Label"],
                           [[x["sheet"], x["direction"], x["text"], x["target_sheet"], x["target_col"], x["label"]] for x in sig])
    links = sig
    refs = OrderedDict()
    for x in links:
        if x["target_sheet"]:
            refs.setdefault(_num(x["sheet"]), set()).add(_num(x["target_sheet"]))
    out["Sheet References"] = (["Sheet", "Refers To Sheets"], [[k, ", ".join(str(i) for i in sorted(v - {k}))] for k, v in sorted(refs.items())])
    return out


def selected(pid, only=None):
    """datasets(pid) restricted to the names in `only` (all when empty), keeping the standard order."""
    ds = datasets(pid)
    return ds if not only else OrderedDict((k, v) for k, v in ds.items() if k in only)


def to_xlsx(pid, project_name="", only=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook(); wb.remove(wb.active)
    names = list(selected(pid, only).keys())
    wb.properties.title = f"{project_name} - {names[0] if len(names) == 1 else 'Engineering Data'}" if project_name else "Engineering Data"
    wb.properties.creator = "PRG Drawing Data Extractor"
    thin = Side(style="thin", color="BBBBBB")
    for name, (cols, rows) in selected(pid, only).items():
        ws = wb.create_sheet(name[:31])
        ws.append([f"{project_name} - {name}"]); ws["A1"].font = Font(bold=True, size=13)
        ws.append(cols)
        for c in ws[2]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=3):
            for c in row:
                c.border = Border(top=thin, bottom=thin, left=thin, right=thin); c.alignment = Alignment(vertical="top", wrap_text=True)
        for i, col in enumerate(cols, 1):
            mx = max([len(str(col))] + [len(str(r[i - 1])) for r in rows[:300]]) if rows else len(col)
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, mx + 2), 50)
        ws.freeze_panes = "A3"; ws.auto_filter.ref = f"A2:{get_column_letter(len(cols))}{max(2, len(rows)+2)}"
    b = io.BytesIO(); wb.save(b); return b.getvalue()


def to_csv(pid, name):
    cols, rows = datasets(pid)[name]
    s = io.StringIO(); w = csv.writer(s); w.writerow(cols); w.writerows(rows)
    return s.getvalue().encode("utf-8-sig")


def to_csv_zip(pid, only=None):
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        for name, (cols, rows) in selected(pid, only).items():
            s = io.StringIO(); w = csv.writer(s); w.writerow(cols); w.writerows(rows)
            z.writestr(re.sub(r"[^A-Za-z0-9]+", "_", name) + ".csv", s.getvalue().encode("utf-8-sig"))
    return b.getvalue()


def to_xml(pid, project_name="", only=None):
    out = [f'<?xml version="1.0" encoding="UTF-8"?>\n<Project name="{escape(project_name)}">']
    for name, (cols, rows) in selected(pid, only).items():
        tag = re.sub(r"[^A-Za-z0-9]+", "", name)
        out.append(f"  <{tag}>")
        keys = [re.sub(r"[^A-Za-z0-9]+", "", c) or "Col" for c in cols]
        for r in rows:
            out.append("    <Row>" + "".join(f"<{k}>{escape(str(v))}</{k}>" for k, v in zip(keys, r)) + "</Row>")
        out.append(f"  </{tag}>")
    out.append("</Project>")
    return "\n".join(out).encode("utf-8")


def to_pdf(pid, project_name="", only=None):
    from reportlab.lib.pagesizes import A3, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    names = list(selected(pid, only).keys())
    what = names[0] if len(names) == 1 else "Engineering Data"
    b = io.BytesIO()
    doc = SimpleDocTemplate(b, pagesize=landscape(A3), leftMargin=20, rightMargin=20, topMargin=20, bottomMargin=20,
                            title=f"{project_name} - {what}" if project_name else what,      # shown as the tab / window title
                            author="PRG Drawing Data Extractor", creator="PRG Drawing Data Extractor",
                            subject=", ".join(names)[:500])
    st = getSampleStyleSheet(); small = st["BodyText"].clone("s"); small.fontSize = 6.5; small.leading = 8
    el = []
    for name, (cols, rows) in selected(pid, only).items():
        el += [Paragraph(f"<b>{escape(project_name)} — {escape(name)}</b>", st["Heading2"]), Spacer(1, 6)]
        if not rows:
            el += [Paragraph("(no data)", st["BodyText"]), PageBreak()]; continue
        data = [[Paragraph(f"<b>{escape(c)}</b>", small) for c in cols]] + [[Paragraph(escape(str(v)), small) for v in r] for r in rows]
        t = Table(data, repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDE8F3")), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                               ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        el += [t, PageBreak()]
    doc.build(el); return b.getvalue()
